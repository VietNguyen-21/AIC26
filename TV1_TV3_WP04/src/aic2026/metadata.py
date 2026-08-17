"""Organizer metadata import, strict video matching, deduplication, and video-level
retrieval."""

from __future__ import annotations

import csv
import json
import math
import re
from datetime import date, datetime
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence
from urllib.parse import parse_qs, urlparse

from .contracts import MediaRecord, MetadataImportReport, MetadataRecord, SearchCandidate
from .ocr import normalize_search_text, strip_vietnamese_diacritics
from .temporal import TemporalRegistry
from .text_index import TextDocument, search_text_index
from .utils import read_jsonl, sha256_file, stable_json_hash, utcnow_iso, write_json, write_jsonl, write_parquet_optional


# Organizer metadata is video-level evidence and must never invent a submission frame.
class MetadataImportError(RuntimeError):
    """Raised when strict organizer metadata import cannot resolve the corpus."""


@dataclass
class MetadataImportResult:
    records: list[MetadataRecord]
    matched_rows: int
    unmatched_rows: int
    invalid_rows: int
    ambiguous_rows: int = 0
    duplicate_rows: int = 0
    source_files: list[Path] = field(default_factory=list)
    unmatched_examples: list[dict[str, Any]] = field(default_factory=list)
    ambiguous_examples: list[dict[str, Any]] = field(default_factory=list)
    duplicate_examples: list[dict[str, Any]] = field(default_factory=list)
    artifact_paths: list[Path] = field(default_factory=list)


def _field(row: dict[str, Any], *names: str) -> Any:
    lowered = {str(key).strip().lower(): value for key, value in row.items()}
    for name in names:
        value = lowered.get(name.lower())
        if value not in (None, ""):
            return value
    return None


def _youtube_id(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    if re.fullmatch(r"[A-Za-z0-9_-]{6,20}", text):
        return text
    try:
        parsed = urlparse(text)
        if parsed.hostname and "youtu" in parsed.hostname:
            if parsed.hostname.endswith("youtu.be"):
                return parsed.path.strip("/").split("/")[0] or None
            query_id = parse_qs(parsed.query).get("v", [None])[0]
            if query_id:
                return query_id
            parts = [part for part in parsed.path.split("/") if part]
            if parts and parts[-1] not in {"watch", "shorts", "embed"}:
                return parts[-1]
            if len(parts) >= 2 and parts[-2] in {"shorts", "embed"}:
                return parts[-1]
    except Exception:
        return None
    return None


def _tags(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, tuple):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return [str(item).strip() for item in parsed if str(item).strip()]
        except json.JSONDecodeError:
            pass
    return [item.strip() for item in re.split(r"[,;|]", text) if item.strip()]


def _iter_json_payload(payload: Any) -> Iterable[dict[str, Any]]:
    if isinstance(payload, list):
        for item in payload:
            if isinstance(item, dict):
                yield item
    elif isinstance(payload, dict):
        for key in ("items", "videos", "records", "data"):
            value = payload.get(key)
            if isinstance(value, list):
                yield from _iter_json_payload(value)
                return
        yield payload


def _read_source_file(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        return [dict(item) for item in read_jsonl(path)]
    if suffix == ".json":
        return list(_iter_json_payload(json.loads(path.read_text(encoding="utf-8-sig"))))
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle, delimiter=delimiter)]
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("Install openpyxl to import organizer Excel metadata") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
            except StopIteration:
                continue
            for values in iterator:
                row = {headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]}
                if any(value not in (None, "") for value in row.values()):
                    rows.append(row)
        return rows
    return []


def metadata_source_fingerprint(root: str | Path, globs: Sequence[str]) -> str:
    base = Path(root)
    files = sorted({path.resolve() for pattern in globs for path in base.glob(pattern) if path.is_file()}) if base.exists() else []
    return stable_json_hash([{"path": path.name, "sha256": sha256_file(path)} for path in files])


def _alias_variants(value: Any) -> list[str]:
    if value in (None, ""):
        return []
    text = str(value).strip()
    if not text:
        return []
    portable = text.replace("\\", "/")
    path = Path(portable)
    candidates = [text, text.lower(), path.name, path.name.lower(), path.stem, path.stem.lower()]
    youtube_id = _youtube_id(text)
    if youtube_id:
        candidates.extend([youtube_id, youtube_id.lower()])
    # Preserve strength/order while removing duplicates.
    return list(dict.fromkeys(candidate for candidate in candidates if candidate))


def _video_lookup(media: Sequence[MediaRecord]) -> dict[str, set[str]]:
    lookup: dict[str, set[str]] = {}
    for item in media:
        path = Path(str(item.original_video_path).replace("\\", "/"))
        aliases = [item.video_id, path.name, path.stem, item.source_sha256]
        for alias in aliases:
            for variant in _alias_variants(alias):
                lookup.setdefault(variant, set()).add(item.video_id)
    return lookup


def _resolve_video_id(
    row: dict[str, Any], lookup: dict[str, set[str]]
) -> tuple[str | None, str | None, dict[str, Any] | None]:
    candidates = [
        ("video_id", _field(row, "video_id", "video id", "id")),
        ("filename", _field(row, "filename", "file_name", "video_file", "video_name", "path")),
        ("youtube_id", _field(row, "youtube_id", "youtube_video_id", "youtube id")),
        ("url", _field(row, "url", "youtube_url", "video_url", "link")),
        ("sha256", _field(row, "sha256", "source_sha256", "checksum")),
    ]
    ambiguous: list[dict[str, Any]] = []
    for matched_by, value in candidates:
        if value in (None, ""):
            continue
        for alias in _alias_variants(value):
            video_ids = sorted(lookup.get(alias, set()))
            if len(video_ids) == 1:
                return video_ids[0], matched_by, None
            if len(video_ids) > 1:
                ambiguous.append(
                    {
                        "matched_by": matched_by,
                        "alias": alias,
                        "candidate_video_ids": video_ids,
                    }
                )
    if ambiguous:
        return None, None, {"matches": ambiguous}
    return None, None, None


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) or math.isinf(value) else value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(item) for item in value]
    return str(value)


def _record_from_row(
    run_id: str,
    video_id: str,
    matched_by: str,
    row: dict[str, Any],
    source_path: Path,
) -> MetadataRecord:
    title = _field(row, "title", "video_title", "name")
    description = _field(row, "description", "desc", "video_description")
    tags = _tags(_field(row, "tags", "keywords", "keyword"))
    channel = _field(row, "channel", "channel_name", "uploader", "author")
    upload_date = _field(row, "upload_date", "published_at", "publish_date", "date")
    language = _field(row, "language", "lang")
    youtube_value = _field(row, "youtube_id", "youtube_video_id", "url", "youtube_url", "video_url")
    youtube_video_id = _youtube_id(youtube_value)
    text_parts = [str(value).strip() for value in (title, description, " ".join(tags), channel) if value not in (None, "")]
    text = " ".join(text_parts)
    safe_row = _json_safe(row)
    row_hash = stable_json_hash(safe_row)
    return MetadataRecord(
        preprocess_run_id=run_id,
        metadata_id=f"organizer_youtube:{video_id}:{row_hash[:24]}",
        video_id=video_id,
        source="organizer_youtube",
        title=None if title is None else str(title),
        description=None if description is None else str(description),
        tags=tags,
        channel=None if channel is None else str(channel),
        upload_date=None if upload_date is None else str(upload_date),
        language=None if language is None else str(language),
        text=text or None,
        normalized_text=normalize_search_text(text) if text else None,
        normalized_text_no_diacritics=normalize_search_text(strip_vietnamese_diacritics(text)) if text else None,
        youtube_video_id=youtube_video_id,
        source_record_sha256=row_hash,
        matched_by=matched_by,
        source_path=str(source_path),
        raw_fields=safe_row,
        created_at_utc=utcnow_iso(),
    )


def import_organizer_youtube_metadata(
    run_id: str,
    run_root: str | Path,
    media: Sequence[MediaRecord],
    config: Any,
) -> MetadataImportResult:
    """Import organizer metadata with strict matching, ambiguity detection, and deduplication."""
    root = Path(run_root)
    source_root = Path(config.organizer_metadata_root)
    patterns = list(config.organizer_metadata_globs)
    files = (
        sorted(
            {
                path.resolve()
                for pattern in patterns
                for path in source_root.glob(pattern)
                if path.is_file()
            }
        )
        if source_root.exists()
        else []
    )
    lookup = _video_lookup(media)
    records: list[MetadataRecord] = []
    matched = unmatched = ambiguous = duplicate = invalid = total_rows = 0
    unmatched_examples: list[dict[str, Any]] = []
    ambiguous_examples: list[dict[str, Any]] = []
    duplicate_examples: list[dict[str, Any]] = []
    seen_source_hashes: dict[tuple[str, str, str], dict[str, Any]] = {}
    example_limit = int(config.max_unmatched_examples)

    for source_file in files:
        try:
            rows = _read_source_file(source_file)
        except Exception as exc:
            invalid += 1
            if len(unmatched_examples) < example_limit:
                unmatched_examples.append(
                    {"source_file": str(source_file), "error": f"{type(exc).__name__}: {exc}"}
                )
            continue
        for index, row in enumerate(rows):
            total_rows += 1
            try:
                video_id, matched_by, ambiguity = _resolve_video_id(row, lookup)
                if video_id is None:
                    if ambiguity is not None:
                        ambiguous += 1
                        if len(ambiguous_examples) < example_limit:
                            ambiguous_examples.append(
                                {
                                    "source_file": str(source_file),
                                    "row_index": index,
                                    "row": _json_safe(row),
                                    **ambiguity,
                                }
                            )
                    else:
                        unmatched += 1
                        if len(unmatched_examples) < example_limit:
                            unmatched_examples.append(
                                {
                                    "source_file": str(source_file),
                                    "row_index": index,
                                    "row": _json_safe(row),
                                }
                            )
                    continue

                record = _record_from_row(
                    run_id, video_id, matched_by or "unknown", row, source_file
                )
                source_hash = record.source_record_sha256 or ""
                dedupe_key = (str(record.source), record.video_id, source_hash)
                first = seen_source_hashes.get(dedupe_key)
                if first is not None:
                    duplicate += 1
                    if len(duplicate_examples) < example_limit:
                        duplicate_examples.append(
                            {
                                "source_file": str(source_file),
                                "row_index": index,
                                "video_id": video_id,
                                "source_record_sha256": source_hash,
                                "first_source_file": first["source_file"],
                                "first_row_index": first["row_index"],
                            }
                        )
                    continue
                seen_source_hashes[dedupe_key] = {
                    "source_file": str(source_file),
                    "row_index": index,
                }
                records.append(record)
                matched += 1
            except Exception as exc:
                invalid += 1
                if len(unmatched_examples) < example_limit:
                    unmatched_examples.append(
                        {
                            "source_file": str(source_file),
                            "row_index": index,
                            "error": f"{type(exc).__name__}: {exc}",
                        }
                    )

    records.sort(key=lambda item: (item.video_id, item.metadata_id))
    source_path = root / "metadata" / "organizer_youtube.jsonl"
    parquet_path = root / "metadata" / "organizer_youtube.parquet"
    report_path = root / "reports" / "metadata_import.json"
    strict_failed = bool(config.strict_unknown_video) and (unmatched > 0 or ambiguous > 0)
    status = "missing_optional_source" if not files else ("failed" if strict_failed else "completed")

    report = MetadataImportReport(
        preprocess_run_id=run_id,
        source_root=str(source_root),
        source_files=[str(path) for path in files],
        source_fingerprint=metadata_source_fingerprint(source_root, patterns),
        total_rows=total_rows,
        matched_rows=matched,
        unmatched_rows=unmatched,
        ambiguous_rows=ambiguous,
        duplicate_rows=duplicate,
        invalid_rows=invalid,
        record_count=len(records),
        strict_unknown_video=bool(config.strict_unknown_video),
        status=status,
        unmatched_examples=unmatched_examples,
        ambiguous_examples=ambiguous_examples,
        duplicate_examples=duplicate_examples,
        created_at_utc=utcnow_iso(),
    )
    write_json(report_path, report.model_dump(mode="json"))

    if strict_failed:
        # Never leave a previous successful organizer artifact looking current
        # after a strict import failure. Rebuild the combined artifact from the
        # remaining provenance sources before surfacing the failure.
        source_path.unlink(missing_ok=True)
        parquet_path.unlink(missing_ok=True)
        consolidate_metadata_artifacts(root)
        raise MetadataImportError(
            "strict_unknown_video rejected organizer metadata: "
            f"unmatched_rows={unmatched}, ambiguous_rows={ambiguous}; "
            f"see {report_path}"
        )

    payload = [item.model_dump(mode="json") for item in records]
    write_jsonl(source_path, payload)
    write_parquet_optional(parquet_path, payload)
    artifacts = [source_path, report_path]
    if parquet_path.exists():
        artifacts.append(parquet_path)
    return MetadataImportResult(
        records=records,
        matched_rows=matched,
        unmatched_rows=unmatched,
        invalid_rows=invalid,
        ambiguous_rows=ambiguous,
        duplicate_rows=duplicate,
        source_files=files,
        unmatched_examples=unmatched_examples,
        ambiguous_examples=ambiguous_examples,
        duplicate_examples=duplicate_examples,
        artifact_paths=artifacts,
    )


def write_technical_metadata(media: Sequence[MediaRecord], run_root: str | Path) -> list[MetadataRecord]:
    root = Path(run_root)
    rows: list[MetadataRecord] = []
    for item in media:
        text = f"video {item.video_id} duration {item.duration_ms / 1000:.1f} seconds resolution {item.width_px}x{item.height_px} codec {item.codec} audio {item.has_audio}"
        rows.append(
            MetadataRecord(
                preprocess_run_id=item.preprocess_run_id,
                metadata_id=f"technical:{item.video_id}",
                video_id=item.video_id,
                source="technical",
                text=text,
                normalized_text=normalize_search_text(text),
                normalized_text_no_diacritics=normalize_search_text(strip_vietnamese_diacritics(text)),
                source_record_sha256=stable_json_hash({"video_id": item.video_id, "source_sha256": item.source_sha256, "text": text}),
                matched_by="video_id",
                created_at_utc=utcnow_iso(),
            )
        )
    payload = [item.model_dump(mode="json") for item in rows]
    write_jsonl(root / "metadata" / "technical.jsonl", payload)
    write_parquet_optional(root / "metadata" / "technical.parquet", payload)
    return rows


def consolidate_metadata_artifacts(run_root: str | Path) -> list[MetadataRecord]:
    """Merge metadata sources without overwriting provenance or duplicating organizer rows."""
    root = Path(run_root)
    records: list[MetadataRecord] = []
    for source_file in (root / "metadata" / "technical.jsonl", root / "metadata" / "organizer_youtube.jsonl", root / "metadata" / "auto_semantic.jsonl", root / "metadata" / "user_annotation.jsonl"):
        for row in read_jsonl(source_file):
            records.append(MetadataRecord.model_validate(row))
    # Preserve provenance sources separately. Organizer records receive an
    # additional source-hash guard so repeated rows cannot bias the text index
    # even when artifacts were produced by an older importer.
    unique: dict[str, MetadataRecord] = {}
    organizer_source_hashes: set[tuple[str, str, str]] = set()
    for item in records:
        if str(item.source) == "organizer_youtube" and item.source_record_sha256:
            key = (str(item.source), item.video_id, item.source_record_sha256)
            if key in organizer_source_hashes:
                continue
            organizer_source_hashes.add(key)
        unique.setdefault(item.metadata_id, item)
    records = sorted(unique.values(), key=lambda item: (item.video_id, str(item.source), item.metadata_id))
    payload = [item.model_dump(mode="json") for item in records]
    write_jsonl(root / "metadata" / "metadata.jsonl", payload)
    write_parquet_optional(root / "metadata" / "metadata.parquet", payload)
    write_json(
        root / "reports" / "metadata_summary.json",
        {
            "record_count": len(records),
            "source_counts": {source: sum(str(item.source) == source for item in records) for source in sorted({str(item.source) for item in records})},
            "video_count": len({item.video_id for item in records}),
            "created_at_utc": utcnow_iso(),
        },
    )
    return records


def build_metadata_documents(run_root: str | Path) -> list[TextDocument]:
    documents: list[TextDocument] = []
    for row in read_jsonl(Path(run_root) / "metadata" / "metadata.jsonl"):
        record = MetadataRecord.model_validate(row)
        searchable = " ".join(value for value in [record.text, record.normalized_text, record.normalized_text_no_diacritics] if value)
        if not searchable:
            continue
        documents.append(TextDocument(record.metadata_id, searchable, {**record.model_dump(mode="json"), "source": "metadata"}))
    return documents


def _representative_frame(run_root: Path, video_id: str) -> tuple[int, int, bool]:
    temporal_path = run_root / "temporal" / "temporal_frames.jsonl"
    rows = [row for row in read_jsonl(temporal_path) if row.get("video_id") == video_id]
    if rows:
        rows.sort(key=lambda row: (int(row.get("timestamp_ms", 0)), int(row.get("frame_id", 0))))
        middle = rows[len(rows) // 2]
        return int(middle["frame_id"]), int(middle["timestamp_ms"]), True
    frame_rows = [row for row in read_jsonl(run_root / "frames.jsonl") if row.get("video_id") == video_id]
    if frame_rows:
        frame_rows.sort(key=lambda row: (int(row.get("timestamp_ms", 0)), int(row.get("frame_id", 0))))
        middle = frame_rows[len(frame_rows) // 2]
        return int(middle["frame_id"]), int(middle["timestamp_ms"]), True
    return 0, 0, False


def metadata_search(
    query_id: str,
    query: str,
    run_id: str,
    run_root: str | Path,
    k: int = 100,
    *,
    settings: Any | None = None,
) -> list[SearchCandidate]:
    """Search video-level metadata while keeping candidates non-submittable until localized."""
    return search_text_index(
        query_id,
        query,
        run_id,
        run_root,
        k,
        settings=settings,
        source_filter={"metadata"},
    )
